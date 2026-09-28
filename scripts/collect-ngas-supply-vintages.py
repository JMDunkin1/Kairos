#!/usr/bin/env python3
"""Retain EIA archived STEO gas supply vintages for a research-only replay.

Uses the latest completed month and its prior-year counterpart within each
archived workbook. No current API series or future forecast month is used.
Run with Python containing openpyxl; output remains ignored beneath .local.
"""
import argparse
import concurrent.futures
import calendar
import datetime as dt
import hashlib
import html
import json
import pathlib
import re
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
PARSER = argparse.ArgumentParser()
PARSER.add_argument('--refresh-current', action='store_true')
PARSER.add_argument('--output-root', type=pathlib.Path)
ARGS = PARSER.parse_args()
OUT = ARGS.output_root or ROOT / ('.local/qore/research/ngas-simplification/forward-supply-vintages' if ARGS.refresh_current else '.local/qore/research/ngas-simplification/supply-vintages')
CUTOFF = dt.datetime.now(dt.timezone.utc).date().isoformat() if ARGS.refresh_current else '2026-09-27'
INDEX_URL = 'https://www.eia.gov/outlooks/steo/outlook.php'
BASE_URL = 'https://www.eia.gov/outlooks/steo/'
FETCHED = dt.datetime.now(dt.timezone.utc).isoformat()


def fetch(url, destination, refresh=False):
    if refresh or not destination.exists():
        with urllib.request.urlopen(url, timeout=60) as response:
            payload = response.read()
        destination.write_bytes(payload)
        versions = OUT / "retained-payloads"
        versions.mkdir(exist_ok=True)
        retained = versions / (hashlib.sha256(payload).hexdigest() + destination.suffix)
        if not retained.exists():
            retained.write_bytes(payload)
    return destination.read_bytes()


def plain(value):
    return html.unescape(re.sub('<[^>]*>', ' ', value)).strip()


def extract(entry):
    filename = entry['url'].rsplit('/', 1)[-1]
    recent_release = entry['releaseDate'] >= (dt.date.fromisoformat(CUTOFF) - dt.timedelta(days=62)).isoformat()
    payload = fetch(entry['url'], OUT / filename, ARGS.refresh_current and recent_release)
    workbook = openpyxl.load_workbook(OUT / filename, read_only=True, data_only=True)
    dates = list(workbook['Dates'].values)
    historical = [row[3] for row in dates if isinstance(row[0], str)
                  and row[0].strip().startswith('Last Historical Month')]
    if len(historical) != 1:
        raise ValueError(f'{filename}: historical cutoff is ambiguous')
    historical = int(historical[0])
    month = f'{historical // 100:04d}-{historical % 100:02d}'
    if month >= entry['releaseDate'][:7]:
        raise ValueError(f'{filename}: historical month not completed at release')
    sheet = workbook['5atab']
    cells = list(sheet.values)
    start_year = int(cells[2][2])
    col = 3 + (historical // 100 - start_year) * 12 + historical % 100 - 1
    prior_col = col - 12
    if prior_col < 3:
        raise ValueError(f'{filename}: missing prior-year comparison')
    if cells[3][col - 1] != calendar.month_abbr[historical % 100]:
        raise ValueError(f'{filename}: unexpected monthly column layout')
    row_map = {row[0]: (index, row) for index, row in enumerate(cells, 1)
               if isinstance(row[0], str)}
    result = dict(entry, file=filename, sha256=hashlib.sha256(payload).hexdigest(),
                  observationMonth=month, previousYearMonth=f'{historical // 100 - 1:04d}-{historical % 100:02d}',
                  units='Bcf/d', vintageType='STEO released estimates, not finalized observed data',
                  availableAfterDate=entry['releaseDate'], cells={})
    for series, name in [('NGPRPUS', 'dryProduction'), ('NGEXPUS_LNG', 'lngExports')]:
        index, row = row_map[series]
        current, previous = row[col - 1], row[prior_col - 1]
        if not all(isinstance(value, (float, int)) and value > 0 for value in [current, previous]):
            raise ValueError(f'{filename}: invalid {series} values')
        result[name + 'BcfPerDay'] = current
        result[name + 'PreviousYearBcfPerDay'] = previous
        result[name + 'YoYChangeBcfPerDay'] = current - previous
        result['cells'][series] = {'current': f'5atab!{get_column_letter(col)}{index}',
                                  'previousYear': f'5atab!{get_column_letter(prior_col)}{index}'}
    result['productionGrowthExceedsLngGrowth'] = result['dryProductionYoYChangeBcfPerDay'] > result['lngExportsYoYChangeBcfPerDay']
    result['supplyGrowthLessLngGrowthBcfPerDay'] = result['dryProductionYoYChangeBcfPerDay'] - result['lngExportsYoYChangeBcfPerDay']
    # Some archived reports are corrected after their original release. Keep
    # notices and delay availability to the latest dated notice conservatively.
    result['notices'] = []
    for url in entry.pop('noticeUrls', []):
        name = url.rsplit('/', 1)[-1]
        notice_payload = fetch(url, OUT / name, ARGS.refresh_current and recent_release)
        notice_text = plain(notice_payload.decode(errors='replace'))
        notice_dates = []
        for match in re.finditer(r'(?:Released|Updated|Revised|Corrected):?\s*([A-Za-z]+\s+\d{1,2},\s*20\d{2})', notice_text):
            try:
                notice_dates.append(dt.datetime.strptime(match[1], '%B %d, %Y').date().isoformat())
            except ValueError:
                pass
        if not notice_dates:
            # Retain an explicit unresolved flag rather than treating an
            # undated correction as known on the original release day.
            result['unresolvedNoticeDate'] = True
        else:
            result['availableAfterDate'] = max(result['availableAfterDate'], *notice_dates)
        result['notices'].append({'url': url, 'file': name,
                                 'sha256': hashlib.sha256(notice_payload).hexdigest(),
                                 'noticeDates': sorted(set(notice_dates))})
    workbook.close()
    return result


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    index_payload = fetch(INDEX_URL, OUT / 'archive-index.html', ARGS.refresh_current)
    entries = []
    for tr in re.findall(r'<tr\b[^>]*>(.*?)</tr>', index_payload.decode(), re.S | re.I):
        cells = re.findall(r'<td\b[^>]*>(.*?)</td>', tr, re.S | re.I)
        if len(cells) < 4:
            continue
        match = re.search(r'href="([^"]+_base\.xlsx)"', cells[3])
        if not match:
            continue
        release = dt.datetime.strptime(plain(cells[1]), '%m/%d/%Y').date().isoformat()
        if not '2020-12-01' <= release <= CUTOFF:
            continue
        entries.append({'releaseDate': release, 'url': BASE_URL + match[1],
                        'noticeUrls': [BASE_URL + url for url in re.findall(r'href="([^"]+)"', ''.join(cells[4:]))]})
    rows, failures = [], []
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(extract, entry): entry for entry in entries}
        for future in concurrent.futures.as_completed(futures):
            try:
                row = future.result()
                rows.append(row)
                print(f"retained {row['releaseDate']} {row['observationMonth']}", flush=True)
            except Exception as error:
                failures.append({'source': futures[future], 'error': str(error)})
    report = {'schemaVersion': 1, 'datasetId': 'eia-steo-released-production-lng-vintages-v1',
              'fetchedAt': FETCHED, 'indexUrl': INDEX_URL,
              'indexSha256': hashlib.sha256(index_payload).hexdigest(),
              'availabilityPolicy': 'Use only if availableAfterDate is strictly before tradeDate; unresolvedNoticeDate rows must not be used.',
              'method': 'Same-vintage last completed month versus same month one year earlier; NGPRPUS and NGEXPUS_LNG in Bcf/d.',
              'caveats': ['STEO last historical month includes contemporaneous estimates rather than final observed production/export data.',
                          'Archived workbooks may incorporate corrections; dated correction notices conservatively delay the entire workbook.',
                          'Production less LNG export growth is a partial supply-demand indicator, not a complete gas balance.'],
              'rows': sorted(rows, key=lambda row: row['releaseDate']), 'failures': failures}
    (OUT / 'supply-vintages.json').write_text(json.dumps(report, indent=2) + '\n')
    observations = [dict(
        releasedAt=row['availableAfterDate'],
        originalReleaseDate=row['releaseDate'],
        month=row['observationMonth'],
        productionBcfd=row['dryProductionBcfPerDay'],
        lngBcfd=row['lngExportsBcfPerDay'],
        priorYearProductionBcfd=row['dryProductionPreviousYearBcfPerDay'],
        priorYearLngBcfd=row['lngExportsPreviousYearBcfPerDay'],
        productionYoYChangeBcfd=row['dryProductionYoYChangeBcfPerDay'],
        lngYoYChangeBcfd=row['lngExportsYoYChangeBcfPerDay'],
        sourceUrl=row['url'], sha256=row['sha256'],
        availability='admit only when releasedAt < target trading session date',
        vintageType=row['vintageType'], cells=row['cells'], notices=row['notices'],
    ) for row in report['rows'] if not row.get('unresolvedNoticeDate')]
    (OUT / 'observations.json').write_text(json.dumps(observations, indent=2) + '\n')
    print(json.dumps({'rows': len(rows), 'failures': failures,
                      'unresolvedNoticeDates': [row['releaseDate'] for row in rows if row.get('unresolvedNoticeDate')]}))
    if failures:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
